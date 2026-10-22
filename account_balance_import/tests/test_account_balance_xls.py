import io

from odoo.tests import HttpCase, tagged
from openpyxl import load_workbook


@tagged("post_install", "-at_install")
class TestAccountBalanceXls(HttpCase):
    def test_template_lists_the_active_company_chart(self):
        """Given a user with two companies, when the template is downloaded for the one active
        in the web client, then it lists that company's chart by code, without the user default
        company accounts nor the payable ones (those go through the partner balances)."""
        other_company = self.env["res.company"].create({"name": "XLS Other Company"})
        self.env.ref("base.user_admin").company_ids |= other_company
        self.authenticate("admin", "admin")
        default_company_account = self.env["account.account"].create(
            {"code": "XLS001", "name": "XLS default", "account_type": "expense", "company_ids": [self.env.company.id]}
        )
        Account = self.env["account.account"].with_company(other_company)
        expense = Account.create(
            {"code": "XLS002", "name": "XLS expense", "account_type": "expense", "company_ids": [other_company.id]}
        )
        payable = Account.create(
            {
                "code": "XLS003",
                "name": "XLS payable",
                "account_type": "liability_payable",
                "company_ids": [other_company.id],
            }
        )

        response = self.url_open(f"/account_balance_import/account_balance_xls?company_id={other_company.id}")

        self.assertEqual(response.status_code, 200)
        rows = list(load_workbook(io.BytesIO(response.content)).active.iter_rows(values_only=True))
        self.assertEqual(rows[0][:3], ("code", "name", "opening_balance"))
        codes = {row[0] for row in rows[1:]}
        self.assertIn(expense.code, codes)
        self.assertNotIn(default_company_account.code, codes)
        self.assertNotIn(payable.code, codes)
